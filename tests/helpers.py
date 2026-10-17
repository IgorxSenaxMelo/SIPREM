"""Gera planilhas sintéticas para os testes (sem depender de dados reais)."""
from pathlib import Path

import numpy as np
import openpyxl

LEVELS = ["Muito baixo", "Baixo", "Moderadamente baixo", "Médio",
          "Moderadamente alto", "Alto", "Muito alto"]
SAATY = ["Igual importância", "Fracamente mais importante", "Moderadamente mais importante",
         "Moderadamente a fortemente mais importante", "Fortemente mais importante",
         "Fortemente a muito fortemente mais importante", "Muito fortemente mais importante",
         "Muito fortemente a extremamente mais importante", "Extremamente mais importante"]
HEAD = ["ID", "Descrição", "Impacto", "Custo", "CAPACIDADE", "Prazo", "RiscoVerif"]


def make_requirements(path, n_req=12, sheets=("E1", "E2", "E3"), extra=("DADOS",), seed=0, bad=None):
    rng = np.random.default_rng(seed)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name in sheets:
        ws = wb.create_sheet(name)
        ws.append(HEAD)
        for i in range(n_req):
            ws.append([f"R{i + 1}", f"Requisito {i + 1}"] + [LEVELS[rng.integers(0, 7)] for _ in range(5)])
    for name in extra:
        wb.create_sheet(name)
    if bad:
        sheet, row, col, value = bad
        wb[sheet].cell(row=row, column=col, value=value)
    wb.save(path)
    return Path(path)


def make_fahp(path, sheets=("E1", "E2", "E3"), extra=("tab",), seed=1):
    rng = np.random.default_rng(seed)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name in sheets:
        ws = wb.create_sheet(name)
        ws.cell(row=4, column=2, value=SAATY[0])
        for c in range(3, 7):
            ws.cell(row=4, column=c, value=SAATY[rng.integers(0, 9)])
    for name in extra:
        wb.create_sheet(name)
    wb.save(path)
    return Path(path)
